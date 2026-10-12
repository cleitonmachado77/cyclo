"""Gera versão externa e enxuta do Relatório de Compliance (para a empresa cliente)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "RELATORIO-COMPLIANCE-IA.docx"


def set_run_font(run, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_para(doc, text, *, size=11, bold=False, italic=False, space_after=6, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
        else:
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        size = 15 if level == 0 else (12 if level == 1 else 11)
        set_run_font(run, size=size, bold=True, color=RGBColor(31, 78, 121))
    return h


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        for part in re.split(r"(\*\*[^*]+\*\*)", item):
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                set_run_font(run, size=10, bold=True)
            else:
                run = p.add_run(part)
                set_run_font(run, size=10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=9, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(cell, "1F4E79")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=9)
            if i % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    add_heading(doc, "Cyclo.ia — Informe de Privacidade e Uso de Dados", 0)
    add_para(
        doc,
        "Documento informativo para Compliance, Jurídico, Segurança da Informação e TI. "
        "Emissão: outubro de 2026. Este texto **não constitui parecer jurídico**.",
        size=9,
        italic=True,
        space_after=10,
    )

    add_heading(doc, "1. Propósito da solução", 1)
    add_para(
        doc,
        "O Cyclo.ia é uma plataforma de **cronometragem operacional (Time & Motion)** que analisa "
        "vídeos de processo enviados pelo usuário e produz tempos e relatórios de movimentos. "
        "A finalidade é medir e melhorar o processo produtivo.",
        size=10,
    )
    add_para(doc, "A solução **não** realiza:", size=10, space_after=2)
    add_bullets(
        doc,
        [
            "Monitoramento contínuo de colaboradores ou da estação de trabalho.",
            "Captura de tela, registro de teclado, áudio do dispositivo ou agente em segundo plano.",
            "Coleta automática de câmera sem ação deliberada do usuário autorizado.",
        ],
    )
    add_para(
        doc,
        "O conteúdo analisado corresponde a **vídeos operacionais encaminhados de forma voluntária** "
        "(por exemplo, registro de posto ou linha de produção), e não a vigilância digital do colaborador.",
        size=10,
    )

    add_heading(doc, "2. Tratamento de dados em linhas gerais", 1)
    add_para(doc, "Em síntese, o fluxo de tratamento observa as seguintes etapas:", size=10, space_after=2)
    add_bullets(
        doc,
        [
            "O usuário autorizado cria a atividade de cronometragem e envia o(s) vídeo(s).",
            "Os arquivos e metadados são armazenados na infraestrutura contratada para a conta do cliente.",
            "Mediante solicitação de análise, o vídeo é processado por provedor de inteligência artificial "
            "especializado em compreensão de vídeo, com o objetivo exclusivo de gerar o relatório de tempos.",
            "Após a análise, a cópia temporária utilizada pelo provedor de IA é eliminada pela solução.",
            "O vídeo original observa prazo de retenção configurável (**padrão de 72 horas**). "
            "O relatório estruturado de tempos permanece disponível para consulta operacional.",
        ],
        numbered=True,
    )

    add_heading(doc, "3. Provedores envolvidos e uso para treinamento", 1)
    add_para(
        doc,
        "A análise utiliza serviços de inteligência artificial de mercado, na qualidade de "
        "**processadores contratados**. Não há treinamento de modelo próprio do Cyclo.ia com os "
        "vídeos do cliente.",
        size=10,
    )
    add_table(
        doc,
        ["Provedor", "Função", "Uso para treinamento (política oficial)"],
        [
            [
                "Google Gemini API (serviço pago)",
                "Análise principal de vídeo",
                "Não utiliza prompts, arquivos (incl. vídeos) nem respostas para aprimorar produtos",
            ],
            [
                "OpenAI API (contingência opcional)",
                "Alternativa em caso de indisponibilidade do provedor principal",
                "Não utiliza dados da API para treinamento, por padrão (salvo opt-in)",
            ],
            [
                "Provedor de armazenamento em nuvem",
                "Guarda de vídeos, metadados e resultados",
                "Não utiliza dados do cliente para treinar modelos sem consentimento escrito",
            ],
        ],
    )
    add_para(
        doc,
        "**Condição essencial:** a proteção contratual da Google quanto à não utilização para "
        "treinamento aplica-se aos **serviços pagos (Paid Services)**. O uso em produção deve "
        "ocorrer com faturamento ativo. Em serviços gratuitos, o provedor pode empregar o conteúdo "
        "para melhoria de produtos — hipótese que deve ser evitida em ambiente corporativo.",
        size=10,
    )
    add_para(
        doc,
        "Mesmo no regime pago, o provedor pode manter registros temporários exclusivamente para "
        "segurança, prevenção de abuso e obrigações legais — **não para treinamento de modelos**. "
        "Organizações com exigência de retenção zero podem avaliar regimes enterprise adicionais "
        "(por exemplo, Zero Data Retention / Vertex AI), mediante contratação específica.",
        size=10,
    )

    add_heading(doc, "4. Garantias e controles relevantes", 1)
    add_bullets(
        doc,
        [
            "Processamento **sob demanda**, iniciado apenas por usuário autorizado.",
            "Eliminação da cópia temporária no provedor de IA após a análise.",
            "Retenção limitada do vídeo (padrão de 72 horas), com limpeza automatizada.",
            "Contingência com segundo provedor de IA **opcional** e desativável, se a política "
            "interna exigir um único processador.",
            "Celebração de Adendo de Processamento de Dados (DPA) com os provedores, conforme "
            "necessidade jurídica da organização.",
        ],
    )

    add_heading(doc, "5. Enquadramento perante políticas internas de “vedação a IA”", 1)
    add_para(
        doc,
        "Se a política institucional veda **qualquer** processamento por modelo de IA, ainda que "
        "como serviço contratado, a solução não atende sem arquitetura alternativa (modelo em "
        "infraestrutura própria). Se o objetivo for impedir que dados corporativos treinem modelos "
        "de terceiros ou bloquear ferramentas consumer na rede, o Cyclo.ia pode ser compatível "
        "mediante conta paga, DPA, controles de acesso e, quando exigido, desativação da contingência "
        "ou regime de retenção zero.",
        size=10,
    )
    add_para(
        doc,
        "A filmagem de postos de trabalho deve observar a LGPD e as normas internas da organização; "
        "a definição da base legal compete à área jurídica do cliente.",
        size=10,
    )

    add_heading(doc, "6. Condições recomendadas para implantação", 1)
    add_bullets(
        doc,
        [
            "Utilizar exclusivamente a API do provedor principal em regime pago.",
            "Formalizar DPA com os processadores envolvidos.",
            "Adotar autenticação corporativa e controle de acesso por perfil/organização.",
            "Manter repositório de vídeos com acesso restrito (sem exposição pública).",
            "Alinhar prazos de retenção de vídeos e relatórios à política interna.",
            "Documentar a ciência dos colaboradores quanto à filmagem para fins de cronometragem.",
        ],
        numbered=True,
    )

    add_heading(doc, "7. Fontes oficiais", 1)
    add_bullets(
        doc,
        [
            "Google — Gemini API Additional Terms: https://ai.google.dev/gemini-api/terms",
            "Google — Zero Data Retention: https://ai.google.dev/gemini-api/docs/zdr",
            "OpenAI — Business data privacy: https://openai.com/business-data/",
            "OpenAI — API data controls: https://developers.openai.com/api/docs/guides/your-data",
        ],
    )

    add_heading(doc, "8. Síntese", 1)
    add_para(
        doc,
        "O Cyclo.ia processa vídeos de processo enviados sob demanda para gerar cronometragem "
        "operacional. Não monitora o desktop do colaborador. Com o provedor principal em conta paga, "
        "a documentação oficial veda o uso de vídeos, prompts e respostas para treinamento ou "
        "aprimoramento de modelos. A inteligência artificial atua como **processador contratado**, "
        "e não como mecanismo de ingestão para treinamento, desde que observadas as condições "
        "descritas neste informe.",
        size=10,
    )
    add_para(
        doc,
        "Políticas de provedores podem ser alteradas; prevalecem o contrato e o DPA celebrados. "
        "Este documento é informativo e não substitui assessoria jurídica especializada.",
        size=9,
        italic=True,
    )

    add_para(
        doc,
        "— Fim do documento —",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        size=9,
        space_after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK: {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
