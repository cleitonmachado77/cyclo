"""Gera a apresentação comercial do Cyclo.ia para cliente."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = 13.333, 7.5

INK = RGBColor(0x00, 0x00, 0x00)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
BG = PAPER
BG_SOFT = PAPER
CARD = PAPER
LINE = INK
TEXT = INK
MUTED = INK
DIM = INK
ACCENT = INK
ACCENT_DEEP = PAPER
GOLD = INK
WHITE = INK

FONT = "Calibri"


def _anchor(tf, anchor):
    mapping = {
        MSO_ANCHOR.TOP: "t",
        MSO_ANCHOR.MIDDLE: "ctr",
        MSO_ANCHOR.BOTTOM: "b",
    }
    tf._txBody.bodyPr.set("anchor", mapping[anchor])


def rect(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    # drop theme shadow
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)
    return shape


def rrect(slide, l, t, w, h, fill, line=None, radius=0.08):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = INK if line is None else line
    shape.line.width = Pt(1.25)
    return shape


def text(
    slide,
    l,
    t,
    w,
    h,
    content,
    size=16,
    bold=False,
    color=TEXT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    spacing=0,
):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    _anchor(tf, anchor)
    lines = content.split("\n") if isinstance(content, str) else content
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(spacing if i else 0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def bg(slide):
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0, 0, W, 0.06, ACCENT)


def footer(slide, page, total, note="Cyclo.ia  ·  Apresentação comercial"):
    rect(slide, 0.55, 7.18, 12.23, 0.01, LINE)
    text(slide, 0.55, 7.22, 8.4, 0.24, note, 11, False, DIM, anchor=MSO_ANCHOR.MIDDLE)
    text(
        slide,
        10.2,
        7.22,
        2.58,
        0.24,
        f"{page}  /  {total}",
        11,
        False,
        DIM,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def kicker(slide, label):
    text(slide, 0.55, 0.28, 8, 0.28, label.upper(), 12, True, ACCENT)


def h1(slide, title, y=0.52):
    text(slide, 0.55, y, 12.2, 0.55, title, 30, True, WHITE)


TOTAL = 14


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Cyclo.ia — Cronometragem operacional com IA"
    prs.core_properties.subject = "Apresentação comercial para cliente"
    prs.core_properties.author = "Cyclo.ia"
    prs.core_properties.category = "Apresentação comercial"

    cover(new_slide(prs))
    agenda(new_slide(prs))
    desafio(new_slide(prs))
    objetivo(new_slide(prs))
    funcionalidades(new_slide(prs))
    fluxo(new_slide(prs))
    medicao(new_slide(prs))
    piloto(new_slide(prs))
    resultados(new_slide(prs))
    leitura(new_slide(prs))
    potencial(new_slide(prs))
    decisoes(new_slide(prs))
    proximos(new_slide(prs))
    fechamento(new_slide(prs))

    out = r"c:\Users\Notebook\Documents\PROJETOS\IA-VC\apresentacoes\Cyclo.ia - Apresentacao para cliente - texto preto.pptx"
    prs.save(out)
    print(out)


def cover(slide):
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0, 0, 0.14, H, ACCENT)
    rect(slide, 0, 6.55, W, 0.95, BG_SOFT)
    text(slide, 0.7, 1.35, 8, 0.32, "CYCLO.IA", 16, True, ACCENT)
    text(
        slide,
        0.7,
        1.8,
        11.5,
        1.7,
        "O tempo do posto,\nmedido e quantificado.",
        48,
        True,
        WHITE,
    )
    text(
        slide,
        0.7,
        3.75,
        10.5,
        0.9,
        "Cronometragem operacional com inteligência artificial.\nDo vídeo da operação ao número que a engenharia consegue comparar, repetir e melhorar.",
        20,
        False,
        MUTED,
    )
    # three facts
    facts = [
        ("01", "Objetivo", "Medir o trabalho real\nem vídeo de posto"),
        ("02", "Entrega", "Tempos, variação\ne perdas separadas"),
        ("03", "Uso", "Padrão, capacidade,\ncusto e melhoria"),
    ]
    for i, (num, title, body) in enumerate(facts):
        x = 0.7 + i * 4.05
        text(slide, x, 6.68, 0.7, 0.55, num, 18, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 0.75, 6.64, 3.0, 0.28, title, 13, True, WHITE)
        text(slide, x + 0.75, 6.9, 3.0, 0.48, body, 12, False, MUTED)


def agenda(slide):
    bg(slide)
    footer(slide, 2, TOTAL)
    kicker(slide, "Roteiro")
    h1(slide, "O que vamos percorrer")
    items = [
        ("01", "O desafio", "Por que o cronômetro no posto limita a decisão."),
        ("02", "Objetivo", "O que o Cyclo.ia se propõe a entregar."),
        ("03", "Funcionalidades", "Como a atividade vira medição e relatório."),
        ("04", "Resultados do piloto", "Números reais de uma montagem filmada."),
        ("05", "Dados quantificados", "O que a empresa passa a conseguir decidir."),
        ("06", "Próximos passos", "Como levar o método para um posto de vocês."),
    ]
    for i, (num, title, body) in enumerate(items):
        col = i % 2
        row = i // 2
        x = 0.55 + col * 6.35
        y = 1.55 + row * 1.7
        rrect(slide, x, y, 6.05, 1.5, CARD, radius=0.06)
        rect(slide, x, y, 0.08, 1.5, ACCENT if row < 2 else GOLD)
        text(slide, x + 0.35, y + 0.28, 1.1, 0.9, num, 26, True, ACCENT if row < 2 else GOLD, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 1.55, y + 0.32, 4.2, 0.4, title, 20, True, WHITE)
        text(slide, x + 1.55, y + 0.78, 4.2, 0.48, body, 15, False, MUTED)


def desafio(slide):
    bg(slide)
    footer(slide, 3, TOTAL)
    kicker(slide, "Contexto")
    h1(slide, "O tempo ainda depende de quem segura o cronômetro")
    text(
        slide,
        0.55,
        1.25,
        12.2,
        0.55,
        "A operação é filmada todos os dias. O que falta é transformar essa observação em amostra, com início, fim e duração que a engenharia possa usar.",
        16,
        False,
        MUTED,
    )
    cards = [
        ("Amostra curta", "Uma visita ao posto observa poucos ciclos. O número representa aquele momento, com pouca repetição."),
        ("Leitura variável", "O instante em que o movimento começa e termina muda de analista para analista."),
        ("Vídeo sem número", "A gravação prova o que aconteceu. A planilha guarda um tempo solto, difícil de auditar."),
        ("Decisão que espera", "Padrão, balanceamento, custo e melhoria ficam parados até o estudo manual fechar."),
    ]
    for i, (title, body) in enumerate(cards):
        x = 0.55 + (i % 4) * 3.15
        y = 2.1
        rrect(slide, x, y, 3.0, 3.15, CARD, radius=0.07)
        rect(slide, x + 0.28, y + 0.32, 0.42, 0.08, ACCENT)
        text(slide, x + 0.28, y + 0.6, 2.45, 0.9, title, 20, True, WHITE)
        text(slide, x + 0.28, y + 1.6, 2.45, 1.25, body, 14, False, MUTED)
    text(
        slide,
        0.55,
        5.5,
        12.2,
        1.4,
        "Sem tempo quantificado, capacidade, custo de mão de obra e ganho de melhoria nascem de estimativa.\nO Cyclo.ia existe para colocar a mesma régua em cada vídeo.",
        18,
        False,
        WHITE,
    )


def objetivo(slide):
    bg(slide)
    footer(slide, 4, TOTAL)
    kicker(slide, "Objetivo")
    h1(slide, "Medir o trabalho real e devolver um número utilizável")
    goals = [
        (
            "01",
            "Medir o que acontece",
            "Cronometrar os movimentos do posto: ciclo, montagem, busca de peça, deslocamento, espera e interrupção. A regra do que conta como ciclo é definida pela engenharia, no prompt da atividade.",
        ),
        (
            "02",
            "Consolidar a amostra",
            "Vários vídeos do mesmo posto entram em fila. Ao final, cada movimento tem quantidade de amostras, média, mínimo e máximo. A variação fica visível, junto com a média.",
        ),
        (
            "03",
            "Sustentar a decisão",
            "O relatório vira base para tempo de referência, leitura de perdas, capacidade do posto, custo do ciclo e comparação antes e depois de um ajuste de método.",
        ),
    ]
    for i, (num, title, body) in enumerate(goals):
        y = 1.5 + i * 1.75
        rrect(slide, 0.55, y, 12.23, 1.6, CARD, radius=0.06)
        text(slide, 0.85, y + 0.28, 1.2, 1.05, num, 28, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 2.3, y + 0.22, 9.9, 0.42, title, 22, True, WHITE)
        text(slide, 2.3, y + 0.72, 9.9, 0.7, body, 15, False, MUTED)


def funcionalidades(slide):
    bg(slide)
    footer(slide, 5, TOTAL)
    kicker(slide, "Funcionalidades")
    h1(slide, "O que o sistema faz")
    items = [
        ("Atividade", "Nome do posto, contexto da operação e o prompt com os movimentos e o critério de ciclo fechado."),
        ("Vídeos em lote", "Vários registros da mesma função. A amostra cresce sem remontar o estudo do zero."),
        ("Esteira", "Análise de um vídeo por vez, com progresso visível, até o lote terminar."),
        ("Cronometragem", "Para cada movimento: início, fim, duração em segundos, status e uma nota do que foi visto."),
        ("Consolidação", "Média, mínimo, máximo e número de amostras por movimento, no relatório final."),
        ("Governança", "Vídeo com retenção padrão de 72 horas. O tempo medido permanece; a gravação não fica indefinida."),
    ]
    for i, (title, body) in enumerate(items):
        col = i % 3
        row = i // 3
        x = 0.55 + col * 4.2
        y = 1.5 + row * 2.65
        rrect(slide, x, y, 4.0, 2.45, CARD, radius=0.07)
        rect(slide, x + 0.28, y + 0.32, 0.46, 0.08, ACCENT if row == 0 else GOLD)
        text(slide, x + 0.28, y + 0.58, 3.45, 0.45, title, 20, True, WHITE)
        text(slide, x + 0.28, y + 1.15, 3.45, 1.05, body, 14, False, MUTED)


def fluxo(slide):
    bg(slide)
    footer(slide, 6, TOTAL)
    kicker(slide, "Operação")
    h1(slide, "Da filmagem ao relatório, em cinco passos")
    steps = [
        ("1", "Definir", "Cadastrar a atividade e escrever o que deve ser reconhecido e o que fecha o ciclo."),
        ("2", "Enviar", "Subir os vídeos do posto. O estudo aceita vários arquivos no mesmo lote."),
        ("3", "Medir", "A IA marca início, fim e duração de cada movimento, vídeo a vídeo."),
        ("4", "Consolidar", "O lote devolve média, mínimo, máximo e a quantidade de amostras."),
        ("5", "Usar", "A engenharia confere no vídeo e aplica o número em padrão, perda e capacidade."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.45 + i * 2.56
        rrect(slide, x, 1.7, 2.42, 4.35, CARD, radius=0.08)
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.78), Inches(2.0), Inches(0.86), Inches(0.86)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = PAPER
        circle.line.color.rgb = INK
        circle.line.width = Pt(1.5)
        text(slide, x + 0.78, 2.0, 0.86, 0.86, num, 22, True, INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 0.16, 3.1, 2.1, 0.5, title, 20, True, WHITE, align=PP_ALIGN.CENTER)
        text(slide, x + 0.16, 3.7, 2.1, 1.9, body, 13, False, MUTED, align=PP_ALIGN.CENTER)
    text(
        slide,
        0.55,
        6.25,
        12.2,
        0.7,
        "A esteira também está preparada para receber vídeo direto de câmera, além do envio manual.",
        15,
        False,
        DIM,
    )


def medicao(slide):
    bg(slide)
    footer(slide, 7, TOTAL)
    kicker(slide, "O que fica quantificado")
    h1(slide, "Cada movimento vira um registro auditável")
    rows = [
        ("Movimento", "Nome estável entre os vídeos: montagem, busca de peças, interrupção."),
        ("Status", "Completo, parcial ou não encontrado. O trabalho visível permanece medido."),
        ("Início e fim", "Timestamp no vídeo, para a engenharia conferir a marcação."),
        ("Duração", "Segundos do trecho. A soma produtiva separa o que gerou ocupação."),
        ("Nota", "Descrição objetiva do que o operador fez naquele intervalo."),
        ("Amostra", "Quantidade, média, mínimo e máximo do mesmo movimento no lote."),
    ]
    for i, (title, body) in enumerate(rows):
        y = 1.42 + i * 0.72
        rrect(slide, 0.55, y, 8.15, 0.64, CARD, radius=0.05)
        text(slide, 0.75, y, 2.15, 0.64, title, 14, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 2.95, y, 5.5, 0.64, body, 13, False, TEXT, anchor=MSO_ANCHOR.MIDDLE)

    rrect(slide, 8.9, 1.42, 3.88, 4.32, CARD, radius=0.07)
    rect(slide, 8.9, 1.42, 3.88, 0.08, GOLD)
    text(slide, 9.15, 1.7, 3.4, 0.9, "Princípio\nde medição", 22, True, WHITE)
    text(
        slide,
        9.15,
        2.75,
        3.4,
        2.7,
        "Se o ciclo oficial não fecha, o sistema registra essa ausência.\n\nO trabalho que aconteceu — montar, buscar, esperar — continua com tempo maior que zero.",
        14,
        False,
        MUTED,
    )
    text(
        slide,
        0.55,
        5.85,
        12.2,
        1.05,
        "É essa separação que torna o dado útil: o tempo produtivo, a perda e o ciclo que ainda não ocorreu ficam em linhas diferentes, com nome estável de um vídeo para o outro.",
        16,
        False,
        WHITE,
    )


def piloto(slide):
    bg(slide)
    footer(slide, 8, TOTAL, "Piloto realizado em 16 de setembro de 2026")
    kicker(slide, "Piloto")
    h1(slide, "Montagem de ferramenta, em posto filmado")
    text(
        slide,
        0.55,
        1.28,
        12.2,
        0.7,
        "A engenharia definiu a regra. O Cyclo.ia aplicou a mesma regra em todos os vídeos do lote.",
        16,
        False,
        MUTED,
    )
    blocks = [
        ("Posto", "Montagem de ferramenta em bancada, com manuseio de componentes."),
        ("Ciclo oficial", "Concluído no instante em que o operador aciona a maçaneta prateada."),
        ("Quem medir", "Um operador, indicado por critério visual no prompt da atividade."),
        ("Amostra", "5 vídeos no lote principal. Uma reanálise de 1 vídeo em atividade separada."),
        ("Processamento", "5 de 5 e 1 de 1 concluídos, sem falha de lote."),
        ("Modelo", "Análise de vídeo com Gemini, em esteira, um arquivo por vez."),
    ]
    for i, (title, body) in enumerate(blocks):
        col = i % 3
        row = i // 3
        x = 0.55 + col * 4.2
        y = 2.15 + row * 2.2
        rrect(slide, x, y, 4.0, 2.0, CARD, radius=0.07)
        text(slide, x + 0.28, y + 0.28, 3.45, 0.4, title, 14, True, ACCENT)
        text(slide, x + 0.28, y + 0.78, 3.45, 0.95, body, 15, False, WHITE)


def resultados(slide):
    bg(slide)
    footer(slide, 9, TOTAL, "Lote principal · 5 vídeos · tempos em segundos")
    kicker(slide, "Resultados")
    h1(slide, "O que cada vídeo devolveu", y=0.48)
    headers = ["Vídeo", "Montagem", "Busca", "Interrupção", "Produtivo", "Leitura"]
    widths = [1.15, 1.45, 1.25, 1.55, 1.45, 5.38]
    data = [
        ["1", "53 s", "5 s", "—", "58 s", "Montagem contínua, com duas buscas curtas de peça."],
        ["2", "56 s", "—", "—", "56 s", "Montagem manual do início ao fim da gravação."],
        ["3", "55 s", "—", "—", "55 s", "Montagem manual, sem acionamento da maçaneta."],
        ["4", "32 s", "4 s", "10 s", "36 s", "Conversa de 10 s separada do tempo produtivo."],
        ["5", "6,4 s", "—", "—", "6,4 s", "Gravação muito curta: só o início do manuseio."],
    ]
    x0, y0 = 0.55, 1.4
    h_head, h_row = 0.48, 0.72
    x = x0
    for w, label in zip(widths, headers):
        rect(slide, x, y0, w - 0.04, h_head, PAPER, line=INK)
        text(slide, x + 0.1, y0, w - 0.2, h_head, label, 13, True, INK, anchor=MSO_ANCHOR.MIDDLE)
        x += w
    for r, row in enumerate(data):
        y = y0 + h_head + 0.06 + r * (h_row + 0.06)
        x = x0
        for c, (w, value) in enumerate(zip(widths, row)):
            rect(slide, x, y, w - 0.04, h_row, PAPER, line=INK)
            text(
                slide,
                x + 0.1,
                y,
                w - 0.2,
                h_row,
                value,
                13 if c == 5 else 15,
                c != 5,
                INK,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            x += w
    text(
        slide,
        0.55,
        6.55,
        12.2,
        0.5,
        "Tempo produtivo soma montagem e busca. A interrupção fica de fora dessa soma. Em nenhum vídeo a maçaneta foi acionada: o ciclo oficial não fechou.",
        13,
        False,
        DIM,
    )


def leitura(slide):
    bg(slide)
    footer(slide, 10, TOTAL, "Consolidação do lote de 5 vídeos · setembro de 2026")
    kicker(slide, "Leitura do piloto")
    h1(slide, "Três movimentos, com média e dispersão", y=0.48)

    kpis = [
        ("8", "trechos de montagem", "média 25,3 s\nmín. 3 s  ·  máx. 56 s"),
        ("3", "buscas de peça", "média 3,0 s\nmín. 2 s  ·  máx. 4 s"),
        ("1", "interrupção", "10 s de conversa\nisolada do produtivo"),
        ("221 s", "tempo classificado", "91% montagem\n4% busca  ·  5% pausa"),
    ]
    for i, (value, label, detail) in enumerate(kpis):
        x = 0.55 + i * 3.2
        rrect(slide, x, 1.35, 3.05, 2.15, CARD, radius=0.07)
        text(slide, x + 0.2, 1.48, 2.65, 0.62, value, 28, True, ACCENT)
        text(slide, x + 0.2, 2.12, 2.65, 0.35, label, 13, True, WHITE)
        text(slide, x + 0.2, 2.5, 2.65, 0.8, detail, 13, False, MUTED)

    text(slide, 0.55, 3.7, 12.2, 0.3, "Composição dos 221 segundos classificados", 14, True, INK)
    text(
        slide,
        0.55,
        4.15,
        12.2,
        0.4,
        "Montagem  ·  202 s  ·  91%          Busca de peças  ·  9 s  ·  4%          Interrupção  ·  10 s  ·  5%",
        16,
        False,
        INK,
    )

    notes = [
        ("O que já dá para afirmar", "A busca é curta e estável, entre 2 e 4 segundos. A pausa de conversa foi identificada e retirada do tempo produtivo. O trabalho de montagem foi medido em todos os vídeos."),
        ("O que este número ainda representa", "A média de 25,3 s é a média dos trechos, porque alguns vídeos têm mais de um segmento. O tempo padrão nasce quando os ciclos fecham pelo critério da maçaneta, em amostra combinada com a engenharia."),
    ]
    for i, (title, body) in enumerate(notes):
        x = 0.55 + i * 6.35
        rrect(slide, x, 5.05, 6.15, 1.95, CARD, radius=0.06)
        text(slide, x + 0.25, 5.18, 5.7, 0.35, title, 14, True, GOLD if i else ACCENT)
        text(slide, x + 0.25, 5.55, 5.7, 1.3, body, 13, False, MUTED)


def potencial(slide):
    bg(slide)
    footer(slide, 11, TOTAL)
    kicker(slide, "Potencial")
    h1(slide, "O que a empresa ganha quando o tempo vira dado")
    text(
        slide,
        0.55,
        1.25,
        12.2,
        0.5,
        "Uma medição isolada informa. Uma base quantificada compara posto, turno, método e o antes e depois de cada mudança.",
        16,
        False,
        MUTED,
    )
    items = [
        ("Tempo de referência", "Média de ciclos fechados, dispersão e folgas. O padrão deixa de ser um número único sem amostra."),
        ("Capacidade do posto", "Ciclos por hora saem do tempo de ciclo. A engenharia vê quanto o posto realmente entrega."),
        ("Custo do ciclo", "Tempo padrão em minutos vezes o custo do minuto de mão de obra. O ciclo ganha valor em reais."),
        ("Perdas visíveis", "Busca, espera, conversa e retrabalho deixam de ser impressão e entram num Pareto, em segundos."),
        ("Balanceamento", "Cada posto comparado com o ritmo da linha. O desvio aparece em tempo, com tamanho."),
        ("Melhoria comprovada", "O mesmo critério, antes e depois do ajuste. O ganho fica em segundos, peças e custo."),
    ]
    for i, (title, body) in enumerate(items):
        col = i % 3
        row = i // 3
        x = 0.55 + col * 4.2
        y = 1.95 + row * 2.4
        rrect(slide, x, y, 4.0, 2.22, CARD, radius=0.07)
        rect(slide, x, y, 4.0, 0.08, ACCENT if row == 0 else GOLD)
        text(slide, x + 0.25, y + 0.28, 3.5, 0.45, title, 18, True, WHITE)
        text(slide, x + 0.25, y + 0.85, 3.5, 1.15, body, 14, False, MUTED)


def decisoes(slide):
    bg(slide)
    footer(slide, 12, TOTAL, "Fórmulas de uso da base · o piloto já alimenta a leitura de perdas")
    kicker(slide, "Da medição à gestão")
    h1(slide, "Quatro contas que o dado quantificado destrava", y=0.48)
    formulas = [
        ("Tempo padrão", "Tempo normal  ×  (1 + folgas)", "A média dos ciclos válidos vira a referência. Mínimo e máximo mostram se o posto está estável."),
        ("Capacidade", "3.600  ÷  tempo de ciclo (s)", "Resultado em ciclos por hora, por posto. Serve para promessa de volume e para o balanceamento."),
        ("Custo", "Tempo padrão (min)  ×  R$/min", "O ciclo passa a ter custo de mão de obra direto, comparável entre métodos."),
        ("Perda", "Soma dos trechos de espera, busca e pausa", "No piloto, 10 s de conversa e 9 s de busca já estão separados dos 202 s de montagem."),
    ]
    for i, (title, formula, body) in enumerate(formulas):
        y = 1.35 + i * 1.35
        rrect(slide, 0.55, y, 12.23, 1.25, CARD, radius=0.05)
        text(slide, 0.8, y + 0.18, 2.5, 0.9, title, 16, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 3.4, y + 0.16, 9.0, 0.4, formula, 18, True, GOLD)
        text(slide, 3.4, y + 0.62, 9.0, 0.48, body, 14, False, MUTED)

    # The last card body might overlap - 4*1.35 = 5.4 + 1.35 start = ends at 1.35+5.4=6.75, plus footer at 7.18. Card height 1.25 so last ends at 1.35+3*1.35+1.25 = 1.35+4.05+1.25 = 6.65. OK.


def proximos(slide):
    bg(slide)
    footer(slide, 13, TOTAL)
    kicker(slide, "Com o cliente")
    h1(slide, "Como começar em um posto de vocês")
    steps = [
        ("1", "Escolher o posto", "Uma função com ciclo reconhecível, volume e interesse de melhoria. Um posto bem definido vale mais do que dez estudos rasos."),
        ("2", "Escrever a régua", "O que inicia o ciclo, o que o encerra e quais perdas importam: busca, espera, conversa, retrabalho."),
        ("3", "Filmar ciclos fechados", "Uma amostra combinada com a engenharia, com o fechamento do ciclo visível. O piloto mostrou o valor disso."),
        ("4", "Validar no vídeo", "Conferir juntos uma parte das marcações. Timestamp contra o que o analista vê. Ajustar o prompt se o nome do movimento variar."),
        ("5", "Publicar a referência", "Congelar média, dispersão e a composição de perdas. Esse é o número que operação, custos e melhoria passam a compartilhar."),
        ("6", "Medir de novo", "Depois de um ajuste de método, repetir com a mesma régua. O ganho aparece em segundos, na mesma base."),
    ]
    for i, (num, title, body) in enumerate(steps):
        col = i % 3
        row = i // 3
        x = 0.55 + col * 4.2
        y = 1.5 + row * 2.65
        rrect(slide, x, y, 4.0, 2.48, CARD, radius=0.07)
        text(slide, x + 0.25, y + 0.22, 0.7, 0.5, num, 22, True, ACCENT)
        text(slide, x + 0.95, y + 0.28, 2.75, 0.45, title, 16, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 0.25, y + 0.95, 3.5, 1.3, body, 13, False, MUTED)


def fechamento(slide):
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0, 0, 0.14, H, ACCENT)
    text(slide, 0.7, 1.45, 8, 0.3, "CYCLO.IA", 16, True, ACCENT)
    text(
        slide,
        0.7,
        1.95,
        11.8,
        1.8,
        "O vídeo mostra o posto.\nO dado quantificado\nmostra o que fazer com ele.",
        36,
        True,
        WHITE,
    )
    text(
        slide,
        0.7,
        4.35,
        11,
        0.9,
        "No piloto, a montagem, a busca e a pausa já têm tempo.\nO próximo passo é fechar ciclos no posto de vocês e transformar essa medição em referência de engenharia.",
        18,
        False,
        MUTED,
    )
    rect(slide, 0.7, 5.55, 2.2, 0.06, GOLD)
    text(
        slide,
        0.7,
        5.8,
        11,
        0.7,
        "Próxima conversa  ·  escolher o posto, a régua do ciclo e o tamanho da amostra.",
        18,
        True,
        WHITE,
    )


if __name__ == "__main__":
    build()
