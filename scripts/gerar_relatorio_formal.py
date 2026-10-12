"""Gera versão formal do Relatório de Compliance em .docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "RELATORIO-COMPLIANCE-IA.docx"


def set_run_font(run, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, color_hex: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_para(doc, text, *, size=11, bold=False, italic=False, space_after=8, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # support **bold** segments
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic)
        else:
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 0 else (13 if level == 1 else 12), bold=True, color=RGBColor(31, 78, 121))
    return h


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r"(\*\*[^*]+\*\*)", item)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                set_run_font(run, size=11, bold=True)
            else:
                run = p.add_run(part)
                set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(cell, "1F4E79")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=9)
            if i % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Capa / identificação
    add_heading(doc, "Relatório de Compliance e Privacidade — Cyclo.ia", 0)
    add_para(
        doc,
        "Documento destinado às áreas de Compliance, Jurídico, Segurança da Informação e Tecnologia da Informação.",
        italic=True,
        size=10,
        space_after=4,
    )
    add_para(doc, "**Produto:** Cyclo.ia — plataforma de cronometragem operacional (Time & Motion) baseada em análise de vídeo.", size=10, space_after=2)
    add_para(doc, "**Data de emissão:** 2 de outubro de 2026.", size=10, space_after=2)
    add_para(
        doc,
        "**Objeto:** descrever o funcionamento técnico da solução, o fluxo de tratamento de dados e as evidências documentais dos provedores quanto à não utilização dos dados do cliente para fins de treinamento de modelos de inteligência artificial, observadas as condições comerciais aplicáveis.",
        size=10,
        space_after=14,
    )

    # 1
    add_heading(doc, "1. Sumário executivo", 1)
    add_para(
        doc,
        "O Cyclo.ia **não constitui ferramenta de monitoramento contínuo de colaboradores**. A solução não realiza captura de tela, não instala agente de software no endpoint, não registra teclado, não captura áudio do dispositivo e não opera em segundo plano no ambiente de trabalho do usuário.",
    )
    add_para(
        doc,
        "O sistema recebe **vídeos operacionais encaminhados de forma deliberada** (por exemplo, registros de posto de trabalho ou linha de produção), processa os movimentos por meio de modelo de inteligência artificial generativa e produz **tempos e relatório de cronometragem**. A finalidade é a medição e a melhoria do processo operacional, e não a vigilância individual do colaborador.",
    )
    add_para(doc, "Quanto ao risco de utilização dos dados para treinamento de modelos de terceiros, destaca-se o seguinte:")

    add_table(
        doc,
        ["Provedor", "Função na solução", "Uso para treinamento (política oficial)", "Condição obrigatória"],
        [
            [
                "Google Gemini API (serviço pago)",
                "Análise principal de vídeo",
                "Não utiliza prompts, arquivos (vídeos/imagens) nem respostas para aprimorar seus produtos",
                "Conta com faturamento ativo (Paid Services)",
            ],
            [
                "OpenAI API",
                "Contingência opcional (quadros/frames)",
                "Não utiliza entradas e saídas da API para treinamento, por padrão",
                "Ausência de opt-in de compartilhamento; preferencialmente Zero Data Retention (ZDR)",
            ],
            [
                "Supabase",
                "Armazenamento de vídeos e resultados",
                "Não utiliza Customer Data para treinar modelos sem consentimento prévio por escrito",
                "Contratação comercial e Termos de Serviço vigentes",
            ],
        ],
    )

    add_para(
        doc,
        "**Conclusão para as áreas de Compliance e Segurança da Informação:** no regime de utilização paga e empresarial padrão das APIs mencionadas, o risco de emprego dos vídeos industriais para treinamento de modelos é **excluído contratualmente** pelas políticas oficiais dos provedores. Tal fato, contudo, **não elimina** o processamento temporário do conteúdo por esses prestadores para a geração da análise, razão pela qual a implantação deve observar controles complementares — conta paga, Adendo de Processamento de Dados (DPA), política de retenção e, quando aplicável, Vertex AI com Zero Data Retention — em conformidade com a política interna da organização.",
    )

    # 2
    add_heading(doc, "2. Escopo funcional da solução", 1)
    add_heading(doc, "2.1 Funcionalidades contempladas", 2)
    add_bullets(
        doc,
        [
            "Criação de atividade de cronometragem, com título, descrição e instruções dos movimentos a serem mensurados.",
            "Envio manual de um ou mais arquivos de vídeo (formatos como mp4, webm e mov).",
            "Submissão do vídeo e das instruções a modelo de compreensão de vídeo (Gemini Agentic).",
            "Extração estruturada de tempos por movimento e geração de relatório consolidado.",
            "Retenção configurável do arquivo de vídeo (padrão de 72 horas) e exclusão automatizada.",
        ],
        numbered=True,
    )

    add_heading(doc, "2.2 Funcionalidades não contempladas", 2)
    add_bullets(
        doc,
        [
            "Captura de tela do computador do colaborador.",
            "Gravação contínua ou agente silencioso no endpoint.",
            "Registro de teclado, monitoramento de área de transferência ou mecanismos equivalentes.",
            "Acionamento automático de câmera sem intervenção do usuário. A fonte “camera” encontra-se prevista em modelo de dados futuro, porém **não está implementada** no fluxo vigente, que admite exclusivamente o envio manual.",
        ],
    )
    add_para(
        doc,
        "Tal delimitação é relevante diante de políticas internas que restrinjam o uso de inteligência artificial ou o monitoramento de colaboradores. O Cyclo.ia classifica-se como **instrumento de análise de processo a partir de mídia voluntariamente fornecida**, e não como ferramenta de observação contínua do ambiente de trabalho digital.",
    )

    # 3
    add_heading(doc, "3. Arquitetura e fluxo de tratamento de dados", 1)
    add_para(
        doc,
        "O fluxo operacional pode ser sintetizado da seguinte forma:",
    )
    add_bullets(
        doc,
        [
            "**Etapa 1 — Origem:** o usuário autorizado (operador ou engenharia industrial), por meio do navegador, cria a atividade, define os movimentos a mensurar e realiza o envio do(s) vídeo(s).",
            "**Etapa 2 — Armazenamento:** o arquivo é depositado no Supabase Storage (repositório activity-videos), com metadados registrados em banco PostgreSQL.",
            "**Etapa 3 — Processamento:** mediante solicitação explícita de análise, o servidor da aplicação (Next.js) obtém cópia temporária do vídeo e o submete à Google Gemini Files API, com processamento em modo Agentic, produzindo estrutura JSON de tempos e movimentos. Ao término, o arquivo é **excluído** no provedor Gemini.",
            "**Etapa 4 — Contingência:** exclusivamente na hipótese de indisponibilidade do Gemini e desde que configurada a chave OPENAI_API_KEY, quadros JPEG extraídos localmente podem ser enviados ao modelo gpt-4o.",
            "**Etapa 5 — Resultado:** o relatório permanece no PostgreSQL da conta Supabase do cliente; o vídeo original observa o prazo de retenção configurado (padrão de 72 horas).",
        ],
        numbered=True,
    )

    add_heading(doc, "3.1 Localização e prazo de cada categoria de dado", 2)
    add_table(
        doc,
        ["Categoria de dado", "Local de tratamento", "Prazo típico"],
        [
            ["Arquivo de vídeo", "Supabase Storage", "Padrão de 72 horas (VIDEO_RETENTION_HOURS), com limpeza automatizada"],
            ["Cópia temporária no servidor", "Diretório temporário do servidor de aplicação", "Exclusivamente durante a análise; removida ao final do processamento"],
            ["Cópia na Gemini Files API", "Infraestrutura Google", "Removida de forma explícita após a análise"],
            ["Quadros enviados à OpenAI (contingência)", "Requisição HTTP à API", "Não persistidos pela aplicação; retenção residual sujeita à política da OpenAI"],
            ["Relatório, tempos e JSON da análise", "PostgreSQL (Supabase)", "Permanece após a expiração do vídeo; sem expiração automática no código atual"],
        ],
    )

    # 4
    add_heading(doc, "4. Modelos de inteligência artificial empregados", 1)
    add_heading(doc, "4.1 Provedor principal: Google Gemini", 2)
    add_table(
        doc,
        ["Item", "Especificação"],
        [
            ["SDK", "@google/genai (GoogleGenAI)"],
            ["Modelo padrão", "gemini-3.8-flash (variável GEMINI_MODEL)"],
            ["Modelos de contingência", "gemini-3.7-flash, gemini-3.6-flash, gemini-3.5-flash"],
            ["Modo de mídia", "MediaProcessing.AGENTIC (Agentic Video Understanding)"],
            ["Entrada", "Arquivo de vídeo completo via Files API e prompt de cronometragem"],
            ["Saída", "Texto/JSON contendo movimentos, tempos e observações"],
        ],
    )
    add_para(
        doc,
        "O modo Agentic permite ao modelo percorrer a linha temporal do vídeo, examinando quadros e segmentos, a fim de reconhecer movimentos operacionais descritos nas instruções. **Não se trata de modelo treinado pelo Cyclo.ia com os vídeos do cliente.**",
    )

    add_heading(doc, "4.2 Provedor de contingência: OpenAI", 2)
    add_table(
        doc,
        ["Item", "Especificação"],
        [
            ["SDK", "openai"],
            ["Modelo padrão", "gpt-4o (variável OPENAI_MODEL)"],
            ["Condição de uso", "Somente se a análise via Gemini falhar e existir OPENAI_API_KEY"],
            ["Entrada", "Até OPENAI_MAX_FRAMES (padrão 24) quadros JPEG e prompt"],
            ["Saída", "Estrutura de análise equivalente (tempos e movimentos)"],
        ],
    )
    add_para(
        doc,
        "Em organizações cuja política veda o envio de dados a múltiplos provedores, a contingência OpenAI pode ser integralmente desabilitada mediante a não configuração da variável OPENAI_API_KEY.",
    )

    # 5
    add_heading(doc, "5. Evidências oficiais quanto à não utilização para treinamento", 1)
    add_para(
        doc,
        "As afirmações a seguir fundamentam-se na **documentação pública vigente dos provedores**, consultada em outubro de 2026. Não se trata de declarações unilaterais do Cyclo.ia, e sim de cláusulas e políticas dos respectivos processadores de dados.",
    )

    add_heading(doc, "5.1 Google Gemini API — serviços pagos (Paid Services)", 2)
    add_para(
        doc,
        "Fonte: Gemini API Additional Terms of Service (https://ai.google.dev/gemini-api/terms), com vigência a partir de 23 de março de 2026 e atualização da página em 28 de abril de 2026.",
        size=10,
        italic=True,
    )
    add_para(doc, "Trecho relevante, em tradução livre e com preservação do sentido jurídico:")
    add_para(
        doc,
        "Quando o cliente utiliza Paid Services (incluindo a cota paga da Gemini API), a Google **não utiliza** seus prompts — compreendendo instruções de sistema, conteúdo em cache e arquivos tais como imagens, vídeos ou documentos — **nem as respostas** para aprimorar seus produtos, processando prompts e respostas conforme o Data Processing Addendum, na qualidade de Data Processor.",
        italic=True,
    )
    add_para(doc, "Complementos oficiais relevantes:")
    add_bullets(
        doc,
        [
            "Nos Unpaid Services (cota gratuita ou Google AI Studio sem faturamento), a Google **pode** utilizar o conteúdo para aprimorar produtos e tecnologias de aprendizado de máquina. Por essa razão, o uso em ambiente corporativo deve ocorrer **exclusivamente** em conta paga.",
            "Nos Paid Services, a Google pode registrar prompts e respostas por período limitado, **unicamente** para detecção de violações da Prohibited Use Policy e cumprimento de obrigações legais — e não para treinamento do modelo.",
            "A documentação de Zero Data Retention (https://ai.google.dev/gemini-api/docs/zdr) reforça a restrição de treinamento nos Paid Services. Para garantia absoluta de retenção zero ou regime contratual mais rigoroso, a Google recomenda o Vertex AI (Gemini Enterprise Agent Platform).",
            "A File API armazena arquivos até exclusão ou expiração; o Cyclo.ia promove a exclusão desses arquivos após a análise (operação ai.files.delete ao término do processamento).",
        ],
    )
    add_para(doc, "Requisitos operacionais para a organização cliente:")
    add_bullets(
        doc,
        [
            "Projeto Google Cloud / Gemini API com faturamento ativo (Paid Service).",
            "Vedação do uso de cota gratuita ou AI Studio unpaid em produção.",
            "Não ativação de compartilhamento voluntário de registros para melhoria de modelo (Data logging and sharing).",
            "Avaliação de migração para Vertex AI com Zero Data Retention, caso a política interna exija retenção zero junto ao provedor.",
        ],
        numbered=True,
    )

    add_heading(doc, "5.2 OpenAI API (contingência)", 2)
    add_para(doc, "Fontes oficiais:")
    add_bullets(
        doc,
        [
            "Business data privacy — https://openai.com/business-data/",
            "Data controls / Your data — https://developers.openai.com/api/docs/guides/your-data",
            "Enterprise privacy — https://openai.com/enterprise-privacy/",
        ],
    )
    add_para(doc, "Pontos oficiais:")
    add_bullets(
        doc,
        [
            "Por padrão, a OpenAI **não treina** modelos com dados da API Platform (entradas e saídas), salvo adesão explícita (opt-in).",
            "A retenção padrão de conteúdo da API para monitoramento de abuso é tipicamente de até 30 dias, conforme DPA; organizações elegíveis podem solicitar Zero Data Retention (ZDR).",
            "O endpoint utilizado pelo Cyclo.ia (chat.completions com imagens) é elegível aos controles de retenção descritos na documentação de data controls.",
        ],
    )
    add_para(doc, "Requisitos operacionais:")
    add_bullets(
        doc,
        [
            "Utilização de conta de API empresarial, e não de planos consumer gratuitos do ChatGPT.",
            "Confirmação, no painel administrativo, da inexistência de opt-in de compartilhamento de dados para treinamento.",
            "Caso a política interna o exija, solicitação de ZDR à OpenAI ou desativação da contingência no Cyclo.ia.",
        ],
        numbered=True,
    )

    add_heading(doc, "5.3 Supabase (armazenamento)", 2)
    add_para(
        doc,
        "Fonte: Supabase Terms of Service (https://supabase.com/terms).",
        size=10,
        italic=True,
    )
    add_para(
        doc,
        "Na medida em que a Supabase emprega recursos de inteligência artificial na prestação de seus serviços, **não utilizará** Customer Data, AI Input ou AI Output para treinar, ajustar (fine-tune) ou aprimorar qualquer modelo de inteligência artificial ou aprendizado de máquina **sem consentimento prévio por escrito** do cliente.",
        italic=True,
    )
    add_para(
        doc,
        "A Supabase atua como processadora dos Customer Data (vídeos e metadados da conta do cliente), dispondo de Adendo de Processamento de Dados (DPA) em https://supabase.com/legal/customer-resources/data-processing-addendum.",
    )

    # 6
    add_heading(doc, "6. Controles implementados na solução Cyclo.ia", 1)
    add_table(
        doc,
        ["Controle", "Implementação", "Evidência técnica"],
        [
            ["Exclusão do vídeo no Gemini após a análise", "ai.files.delete ao término do processamento", "src/lib/gemini.ts"],
            ["Exclusão de arquivos temporários no servidor", "Remoção no bloco de finalização", "src/lib/gemini.ts"],
            ["Retenção de vídeo no armazenamento", "VIDEO_RETENTION_HOURS (padrão 72) e campo expires_at", "src/lib/pipeline.ts; src/lib/store.ts"],
            ["Limpeza automatizada", "Endpoint POST /api/cron/cleanup-videos", "src/app/api/cron/cleanup-videos/route.ts"],
            ["Contingência OpenAI opcional", "Ativada somente se OPENAI_API_KEY estiver configurada", "src/lib/gemini.ts"],
            ["Processamento sob demanda", "Análise iniciada apenas mediante ação do usuário", "Fluxo /analyze"],
            ["Ausência de treinamento próprio", "Inexistência de pipeline de fine-tuning com vídeos do cliente", "Estruturas legadas de training removidas"],
        ],
    )

    # 7
    add_heading(doc, "7. Matriz de riscos e medidas de mitigação", 1)
    add_table(
        doc,
        ["Risco percebido", "Realidade técnica", "Mitigação recomendada"],
        [
            [
                "Utilização dos vídeos para treinamento de modelos",
                "Nas políticas oficiais do Gemini pago e da OpenAI API, o uso para treinamento é vedado por padrão",
                "Conta paga, ausência de opt-in de compartilhamento e celebração de DPA",
            ],
            [
                "Monitoramento integral da estação de trabalho do colaborador",
                "Não há captura de desktop; admite-se apenas o envio de vídeo de processo",
                "Comunicação clara do escopo e política de filmagem operacional com ciência do colaborador (LGPD)",
            ],
            [
                "Permanência indefinida dos dados na infraestrutura do provedor de IA",
                "O arquivo no Gemini é excluído após a análise; podem existir registros limitados de monitoramento de abuso",
                "Conta paga; avaliação de Vertex AI / ZDR; desativação da OpenAI, se necessário",
            ],
            [
                "Exposição pública de vídeo industrial",
                "O repositório activity-videos encontra-se, no schema atual, com leitura pública",
                "Tornar o repositório privado e adotar URLs assinadas antes da entrada em produção corporativa",
            ],
            [
                "Acesso indevido às interfaces da aplicação",
                "A versão atual não possui autenticação de usuário final nas rotas",
                "Implementar autenticação corporativa (SSO/SAML) e isolamento por organização",
            ],
            [
                "Permanência indefinida de resultados sensíveis",
                "O relatório estruturado permanece no PostgreSQL após a expiração do vídeo",
                "Estabelecer política de retenção também para os registros de análise",
            ],
            [
                "Presença de inteligência artificial no ambiente corporativo",
                "O modelo não é executado localmente; o processamento ocorre nos data centers do provedor",
                "Decidir entre: (a) API paga com DPA; (b) Vertex AI em região contratada; ou (c) não utilização da solução",
            ],
        ],
    )

    add_heading(doc, "7.1 Esclarecimento quanto a políticas de vedação absoluta ao uso de IA", 2)
    add_para(
        doc,
        "Caso a política institucional vede, de forma literal, **qualquer** processamento por modelo de inteligência artificial — ainda que na qualidade de processador contratado —, o Cyclo.ia **não atenderá** ao requisito sem alteração de arquitetura (por exemplo, modelo hospedado integralmente em infraestrutura própria).",
    )
    add_para(
        doc,
        "Caso a política vise impedir que dados corporativos sejam empregados no treinamento de modelos de terceiros, ou impedir o uso de ferramentas consumer na rede corporativa, o Cyclo.ia **pode revelar-se compatível**, desde que observadas as seguintes condições:",
    )
    add_bullets(
        doc,
        [
            "Utilização exclusiva da Gemini API em regime pago, ou do Vertex AI.",
            "Desativação da OpenAI ou contratação de Zero Data Retention.",
            "Controles de acesso, repositório privado e celebração dos respectivos DPA e contratos.",
            "Definição, pela área jurídica do cliente, da base legal adequada à filmagem de postos de trabalho à luz da LGPD (interesse legítimo, consentimento ou outra hipótese cabível).",
        ],
        numbered=True,
    )

    # 8
    add_heading(doc, "8. Recomendações para implantação em ambiente de elevada exigência de compliance", 1)
    add_heading(doc, "8.1 Medidas obrigatórias antes da entrada em produção", 2)
    add_bullets(
        doc,
        [
            "Gemini com faturamento ativo (Paid Services). Sem essa condição, a proteção contratual de não utilização para treinamento **não se aplica**.",
            "Repositório de vídeos privado, com URLs assinadas (correção do atributo de leitura pública atualmente presente no schema).",
            "Autenticação corporativa (SSO) e isolamento por organização.",
            "Celebração de DPA com a Google e, se utilizados, com OpenAI e Supabase.",
            "Desabilitação da contingência OpenAI, caso a política admita um único processador.",
            "Registro da ciência dos colaboradores quanto à filmagem do posto para fins de cronometragem, em conformidade com a LGPD e normas internas.",
            "Ativação da rotina de limpeza com segredo de autenticação (CRON_SECRET) e retenção alinhada à política do cliente.",
        ],
        numbered=True,
    )
    add_heading(doc, "8.2 Medidas fortemente recomendadas", 2)
    add_bullets(
        doc,
        [
            "Migração da análise para Vertex AI com Zero Data Retention, se a organização exigir retenção zero junto ao provedor.",
            "Política de retenção também para relatórios estruturados, e não apenas para o arquivo de vídeo.",
            "Registro das Operações de Tratamento (ROPA), com indicação de Google, OpenAI e Supabase como operadores ou processadores, conforme o caso.",
            "Alinhamento da região de processamento (data residency) ao contrato de nuvem do cliente.",
        ],
        numbered=True,
    )

    # 9
    add_heading(doc, "9. Fontes oficiais consultadas", 1)
    add_table(
        doc,
        ["Documento", "Endereço eletrônico"],
        [
            ["Gemini API Additional Terms", "https://ai.google.dev/gemini-api/terms"],
            ["Gemini API Zero Data Retention", "https://ai.google.dev/gemini-api/docs/zdr"],
            ["Gemini API Data logging and sharing", "https://ai.google.dev/gemini-api/docs/logs-policy"],
            ["Gemini API Pricing", "https://ai.google.dev/gemini-api/docs/pricing"],
            ["OpenAI Business data privacy", "https://openai.com/business-data/"],
            ["OpenAI API data controls", "https://developers.openai.com/api/docs/guides/your-data"],
            ["OpenAI Enterprise privacy", "https://openai.com/enterprise-privacy/"],
            ["Supabase Terms of Service", "https://supabase.com/terms"],
            ["Supabase Data Processing Addendum", "https://supabase.com/legal/customer-resources/data-processing-addendum"],
        ],
    )

    # 10
    add_heading(doc, "10. Limitações e ressalvas", 1)
    add_bullets(
        doc,
        [
            "Este documento descreve o comportamento da solução Cyclo.ia e as políticas públicas dos provedores na data de emissão indicada.",
            "As políticas dos provedores podem ser alteradas; prevalecem o contrato e o DPA efetivamente celebrados.",
            "O presente relatório **não constitui parecer jurídico**. Compete à área jurídica do cliente validar o enquadramento à LGPD, as bases legais da filmagem e as cláusulas contratuais definitivas.",
            "Controles ainda pendentes na solução (autenticação e repositório privado) devem ser tratados como **pré-requisitos de implantação**, e não como garantias já vigentes em produção.",
        ],
    )

    # 11
    add_heading(doc, "11. Síntese para o comitê de compliance", 1)
    add_para(
        doc,
        "O Cyclo.ia analisa vídeos de processo enviados sob demanda, com a finalidade de produzir cronometragem operacional. A solução **não monitora** a estação de trabalho do colaborador. Com a Gemini API em conta paga, a Google declara, em seus termos oficiais, que **não utiliza** vídeos, prompts nem respostas para treinar ou aprimorar seus modelos; a OpenAI API e a Supabase apresentam compromissos equivalentes no regime padrão aplicável. O vídeo é excluído junto ao provedor de inteligência artificial após a análise e retido no armazenamento do cliente por prazo limitado (padrão de 72 horas). O emprego da inteligência artificial ocorre na qualidade de **processador contratado**, e não como mecanismo de ingestão para treinamento — desde que a implantação observe as condições e os controles descritos neste relatório.",
    )

    add_para(
        doc,
        "— Fim do documento —",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        size=10,
        space_after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK: {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
