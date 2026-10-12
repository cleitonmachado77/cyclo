"""Converte docs/RELATORIO-COMPLIANCE-IA.md em .docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "RELATORIO-COMPLIANCE-IA.md"
OUT_PATH = ROOT / "docs" / "RELATORIO-COMPLIANCE-IA.docx"


def set_cell_shading(cell, color_hex: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_runs_with_markup(paragraph, content: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", cell_text.strip())
            run = p.add_run(clean)
            run.font.size = Pt(9)
            run.font.name = "Calibri"
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, "1F4E79")
            elif i % 2 == 0:
                set_cell_shading(cell, "F2F2F2")
    doc.add_paragraph()


def main() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    lines = text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                p.paragraph_format.left_indent = Cm(0.5)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if "|" in line and line.strip().startswith("|"):
            if re.match(r"^\|[\s\-:|]+\|$", line.strip()):
                i += 1
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                add_table(doc, table_buf)
                table_buf = []
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_runs_with_markup(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_markup(p, line[2:].strip())
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs_with_markup(p, m.group(2))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs_with_markup(p, line.strip())
        i += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"OK: {OUT_PATH}")
    print(f"Size: {OUT_PATH.stat().st_size} bytes")


if __name__ == "__main__":
    main()
